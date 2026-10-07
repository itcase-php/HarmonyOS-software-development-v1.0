"""Independent parse of packages produced by the modeled SDK flow, not real PDF extraction."""
from pathlib import Path
import json
import zipfile
from io import BytesIO
from lxml import etree
from PIL import Image
from pptx import Presentation
from docx import Document

root = Path(__file__).resolve().parents[1]
fixtures = root / 'tmp/offline-engines/pdf-office-tests'
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'v': 'urn:schemas-microsoft-com:vml'}
for extension in ['docx', 'pptx']:
    source = fixtures / ('rebuilt.' + extension)
    with zipfile.ZipFile(source) as package:
        assert package.testzip() is None
        images = [name for name in package.namelist() if name.endswith('.png')]
        assert len(images) == 2
        for name in images:
            Image.open(BytesIO(package.read(name))).verify()
        if extension == 'docx':
            doc = Document(source)
            assert len(doc.sections) == 2
            assert all(section.page_width.pt == 612 and section.page_height.pt == 792 for section in doc.sections)
            tree = etree.fromstring(package.read('word/document.xml'))
            assert tree.xpath('//w:txbxContent//w:t/text()', namespaces=ns) == ['Editable 12345'] * 2
            shapes = tree.xpath('//v:shape', namespaces=ns)
            assert 'margin-top:292pt' in shapes[0].get('style')  # 792 - 400 - 100
            assert 'margin-top:72pt' in shapes[1].get('style')  # 792 - 720
            assert len(shapes[0].xpath('./v:imagedata', namespaces=ns)) == 1
            assert len(shapes[1].xpath('./v:textbox', namespaces=ns)) == 1
        else:
            deck = Presentation(source)
            assert len(deck.slides) == 2
            assert deck.slide_width.pt == 612 and deck.slide_height.pt == 792
            for slide in deck.slides:
                picture, text = slide.shapes
                assert picture.top.pt == 292 and picture.left.pt == 40
                assert text.top.pt == 72 and text.left.pt == 40
                assert text.text == 'Editable 12345'
                assert text.text_frame.paragraphs[0].runs[0].font.name == 'DejaVu Sans'
                assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 20
report = {'scope': 'PDF to Office code with modeled SDK objects and independent output parsing',
          'result': 'passed', 'targets': ['docx', 'pptx'], 'flowCases': 2, 'failureAndControlCases': 21,
          'objectOrderAndCoordinates': 'passed for modeled bottom-left PDF coordinates',
          'protectionAdmission': 'requires trusted completed input probe; main PDF protection probe not implemented',
          'encryptedAndUnsupportedContent': 'rejected in modeled cases', 'fontSubstitution': 'none selected',
          'actualPdfKitExecution': 'not run', 'layoutFidelity': 'not verified', 'mainRouteIntegration': 'not implemented'}
(root / 'tests/generated/pdf-office-code-validation.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8', newline='\n')
print(json.dumps(report, indent=2))
