"""Independent OOXML parse/editability checks; does not validate page rendering."""
from pathlib import Path
from io import BytesIO
import posixpath
import json
import zipfile
from lxml import etree
from PIL import Image
from docx import Document
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
fixtures = root / 'tmp/offline-engines/package-tests'
expected = '中文 😀 <&> 12345'
namespaces = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
results = []
for extension in ['docx', 'pptx']:
    path = fixtures / ('editable.' + extension)
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        assert len(package.namelist()) == len(set(package.namelist()))
        trees = {name: etree.fromstring(package.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
                 for name in package.namelist() if name.endswith(('.xml', '.rels'))}
        for name, tree in trees.items():
            if not name.endswith('.rels'):
                continue
            base = '' if name == '_rels/.rels' else name.rsplit('/_rels/', 1)[0]
            for relationship in tree:
                assert relationship.get('TargetMode') != 'External'
                target = posixpath.normpath(posixpath.join(base, relationship.get('Target')))
                assert target in package.namelist(), (name, target)
        images = [name for name in package.namelist() if name.endswith('.png')]
        assert len(images) == 1
        Image.open(BytesIO(package.read(images[0]))).verify()
        if extension == 'docx':
            document = Document(path)
            assert len(document.sections) == 2
            tree = trees['word/document.xml']
            assert len(tree.xpath('//w:txbxContent', namespaces=namespaces)) == 2
            assert len(tree.xpath('//w:br', namespaces=namespaces)) == 2
            values = tree.xpath('//w:txbxContent//w:t/text()', namespaces=namespaces)
            assert values == [expected, 'Second line'] * 2
            assert all(section.page_width.pt == 612 and section.page_height.pt == 792 for section in document.sections)
        else:
            deck = Presentation(path)
            assert len(deck.slides) == 2
            assert deck.slide_width.pt == 612 and deck.slide_height.pt == 792
            assert [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame] == [expected + '\vSecond line'] * 2
            # Editing actual text nodes and reopening proves the fixture contains text shapes.
            deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = 'Edited 98765'
            edited = fixtures / 'edited.pptx'
            deck.save(edited)
            assert 'Edited 98765' in Presentation(edited).slides[0].shapes[0].text
    results.append({'format': extension, 'zipCrc': 'passed', 'xmlAndRelationships': 'passed',
                    'editableTextObjects': 2, 'separateImages': 1, 'pagesOrSlides': 2})
report = {'scope': 'OOXML host structure and editability only', 'result': 'passed', 'files': results,
          'deviceReadback': 'pending', 'renderAndLayoutFidelity': 'pending'}
(root / 'tests/generated/office-package-validation.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8', newline='\n')
print(json.dumps(report, indent=2))
